# -*- coding: utf-8 -*-
"""Build the digital human research note DOCX."""

import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = r"D:\iwen-codex\codex-2\deliverables\郭柱-数字人交互调研笔记.docx"

HEAD_FONT = "微软雅黑"
BODY_FONT = "等线"
ACCENT = "1F3864"        # dark blue, table header fill
ROW_TINT = "F2F5FA"      # pale blue, alternate body rows
BORDER = "D9D9D9"


# ---------------------------------------------------------------- primitives

def set_run(run, size=10.5, bold=False, font=BODY_FONT, color="000000"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)
    rfonts.set(qn("w:eastAsia"), font)
    return run


def style_font(style, font, size, bold=False):
    style.font.name = font
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string("000000")
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)
    rfonts.set(qn("w:eastAsia"), font)


def setup_styles(doc):
    style_font(doc.styles["Normal"], BODY_FONT, 10.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(8)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.4
    style_font(doc.styles["Title"], HEAD_FONT, 21, bold=True)
    # The stock Title style carries a bottom border; the document must not show
    # a rule under the title block.
    title_ppr = doc.styles["Title"].element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)
    for name, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)):
        style_font(doc.styles[name], HEAD_FONT, size, bold=True)
    for name, size in (("List Bullet", 10.5), ("List Number", 10.5)):
        style_font(doc.styles[name], BODY_FONT, size)


def para(doc, text, lead=None, size=10.5, after=8, line=1.4, align=None, font=BODY_FONT):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if lead:
        set_run(p.add_run(lead), size=size, bold=True, font=font)
    set_run(p.add_run(text), size=size, font=font)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style="Heading %d" % level)
    pf = p.paragraph_format
    pf.space_before = Pt(16 if level == 1 else 12)
    pf.space_after = Pt(7 if level == 1 else 5)
    pf.line_spacing = 1.2
    pf.keep_with_next = True
    set_run(p.add_run(text), size=15 if level == 1 else 12.5, bold=True, font=HEAD_FONT)
    return p


def bullet(doc, text, lead=None):
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.space_after = Pt(4)
    pf.line_spacing = 1.35
    if lead:
        set_run(p.add_run(lead), size=10.5, bold=True)
    set_run(p.add_run(text), size=10.5)
    return p


# --------------------------------------------------------------------- table

def _borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), BORDER)
        borders.append(el)
    tbl_pr.append(borders)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    margins = OxmlElement("w:tblCellMar")
    for tag, value in (("top", 90), ("left", 120), ("bottom", 90), ("right", 120)):
        el = OxmlElement("w:" + tag)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        margins.append(el)
    tbl_pr.append(margins)


def _shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def _fill_cell(cell, text, size, bold, color, align):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(1)
    pf.space_before = Pt(1)
    pf.line_spacing = 1.25
    set_run(p.add_run(text), size=size, bold=bold, color=color)


def table(doc, headers, rows, widths, aligns=None, size=9,
          keep_together=True, bold_rows=()):
    aligns = aligns or [WD_ALIGN_PARAGRAPH.LEFT] * len(headers)
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.autofit = False
    _borders(t)

    hdr = t.rows[0]
    for idx, text in enumerate(headers):
        _fill_cell(hdr.cells[idx], text, size, True, "FFFFFF", WD_ALIGN_PARAGRAPH.CENTER)
        _shade(hdr.cells[idx], ACCENT)
    _repeat_header(hdr)

    for r_idx, row in enumerate(rows):
        cells = t.add_row().cells
        for c_idx, text in enumerate(row):
            _fill_cell(cells[c_idx], text, size, r_idx in bold_rows, "000000", aligns[c_idx])
            if r_idx % 2 == 1:
                _shade(cells[c_idx], ROW_TINT)

    # Never let a single row break across pages; optionally keep the whole
    # table on one page by chaining the paragraphs forward.
    for r_idx, row in enumerate(t.rows):
        _cant_split(row)
        if keep_together and r_idx < len(t.rows) - 1:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True

    for row in t.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Cm(widths[idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for idx, width in enumerate(widths):
        for row in t.rows:
            row.cells[idx].width = Cm(width)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def footer_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_run(run, size=9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


# ------------------------------------------------------------------- content

def build():
    doc = Document()
    setup_styles(doc)

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    footer_page_number(section)

    # Title block
    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(6)
    set_run(title.add_run("数字人交互模块技术调研与演示方案"), size=21, bold=True, font=HEAD_FONT)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(18)
    set_run(meta.add_run("调研人 郭柱　　日期 2026 年 9 月 25 日　　阶段 调研阶段 国庆前定稿"),
            size=9.5, font=BODY_FONT)

    para(doc, "这份笔记交付给钟老师和团队其他成员，内容是我负责的数字人与前端交互模块的调研结论和演示方案。"
              "调研范围包括开源数字人框架、语音合成与识别方案、口型同步技术，以及与后端对话模型的联动方式。"
              "文中所有项目的星标数、许可证和最近更新时间，来自 2026 年 9 月 25 日检索的公开仓库信息。")
    para(doc, "调研得到的最重要结论是，这个模块不影响初赛成绩。初赛由平台自动运行我们提交的 Docker 镜像，"
              "对对话模型的输出做离线测评，数字人形象和前端界面都不参与打分，它的价值集中在十二月下旬的总决赛答辩与作品展示。")
    para(doc, "因此本模块的选型原则是够用、稳定、不占用训练算力。我推荐以二维卡通形象配合流式语音合成和 WebRTC 推流，"
              "底座直接复用成熟的开源框架，把精力集中在情感联动和现场演示的可靠性上。")

    # 1
    heading(doc, "结论摘要")
    para(doc, "初赛是平台自动运行 Docker 镜像、对对话模型的输出做离线测评，数字人形象与前端界面都不参与打分，"
              "它的价值集中在总决赛答辩与作品展示，所以这一模块的原则是够用、稳定、不占用训练算力。",
         lead="本模块不进入初赛评分链路。")
    para(doc, "推荐二维卡通形象、流式语音合成、WebRTC 推流三条技术组合，底座优先复用 LiveTalking、"
              "Open-LLM-VTuber 或 OpenTalking，不自研渲染与口型算法。",
         lead="推荐的技术路线明确。")
    para(doc, "对话模型输出的情绪标签同时驱动语音合成的情感参数和数字人的表情动作，让共情被看得见。"
              "这是本项目相对通用数字人方案的差异点，也是答辩现场最容易说服评委的部分。",
         lead="创新点在情感联动。")
    para(doc, "首包三秒以内算合格，一点五秒以内算优秀。同行项目公开的实测锚点是级联方案约八吉字节显存、首包约三秒。",
         lead="延迟基准可以量化。")
    para(doc, "初赛镜像只保留对话模型推理，数字人演示独立成服务，避免增加镜像运行失败的风险。",
         lead="数字人服务与初赛镜像分离。")

    # 2
    heading(doc, "模块定位与边界")
    heading(doc, "职责范围", 2)
    para(doc, "我负责数字人形象与渲染方案、口型同步方案、语音输入输出链路、前端交互界面、"
              "与对话模型服务之间的对接协议，以及最终的演示设计。对话模型的训练与优化、数据集与评估口径、"
              "用户需求与竞品调研分别由张力文、唐应杰、王如怡负责，我不在这三块重复投入。")
    para(doc, "我需要向其他模块输出三样东西，分别是对接对话模型的服务接口约定、团队统一使用的情绪标签体系、"
              "以及决赛演示脚本。这三项都是跨模块接口，越早定稿越省事。")

    heading(doc, "数字人服务与初赛镜像分离", 2)
    para(doc, "赛事要求初赛作品是 Docker 镜像，入口为 /participant/run_all.sh，"
              "平台在离线环境中自动读取测试数据、运行推理并按格式输出预测结果，全程无需人工干预。"
              "这条约束推到本模块有两个结论：数字人做得好坏对初赛排名没有影响；"
              "而一旦把语音合成或推流服务放进镜像，启动失败或超时都会让整个镜像判定为运行失败，"
              "赔上全队仅有的五次提交机会。")
    para(doc, "因此我建议数字人演示独立成一个仓库和服务，初赛镜像只保留对话模型推理。"
              "这条判断会影响邹艾融的整合方式，需要在与钟老师讨论时明确下来。")

    heading(doc, "与其他模块的接口", 2)
    table(doc,
          ["对接人", "需要约定的事", "我提供的内容"],
          [
              ["张力文　核心对话模型",
               "推理服务协议，输入多轮历史与当前用户输入，输出流式文本与情绪标签",
               "需要返回的字段清单与流式返回方式"],
              ["唐应杰　数据集与评估",
               "情绪标签的类别要与数据标注和评估指标对齐",
               "情绪标签体系草案"],
              ["王如怡　需求与竞品",
               "目标用户决定数字人的人设与交互风格",
               "形象与交互风格建议"],
              ["邹艾融　统筹",
               "数字人与初赛镜像是否分离，决赛演示采用现场还是录屏",
               "演示技术方案与降级预案"],
          ],
          [3.6, 6.4, 6.2])

    # 3
    heading(doc, "数字人形象的技术路线")
    para(doc, "数字人形象有三条现实可行的路线，差异主要在口型实现方式、算力占用和现场稳定性上。"
              "下表按统一维度把三条路线放在一起比较。")
    table(doc,
          ["维度", "二维卡通形象", "二维真人视频驱动", "三维模型"],
          [
              ["代表方案", "Open-LLM-VTuber、live2d-py", "Wav2Lip、MuseTalk、LatentSync、EchoMimicV2",
               "TalkingHead、Unity 加 uLipSync"],
              ["口型实现", "音频转表情参数，控制口型开合与眨眼", "音频逐帧重绘嘴部像素",
               "音频转 viseme 或表情系数"],
              ["显存占用", "几乎为零", "需要显卡常驻，数吉字节起", "服务端零占用，浏览器渲染"],
              ["观感", "亲和稳定，风格偏虚拟伴侣", "最真实，但容易出现嘴部模糊与抖动", "中等，风格化"],
              ["制作成本", "低，替换模型文件即可", "高，需要形象素材与调参", "中，需要模型与绑定"],
              ["与情感陪伴场景的契合度", "高", "中，真人形象偏客服与主播", "中"],
              ["现场演示风险", "低", "中高，受显存与掉帧影响", "低"],
          ],
          [3.0, 4.6, 4.8, 3.8],
          aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT,
                  WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])
    para(doc, "二维卡通形象加 Live2D 是首选，二维真人视频驱动作为备选。赛题的关键词是数字人情感陪伴与智能对话，"
              "陪伴型产品的主流形态本来就是二次元形象，用户接受度高。更实际的原因是，"
              "二维路线把显卡全部留给对话模型和语音合成，对我们只有两卡 4090 的开发实例非常关键，"
              "而且决赛现场演示最怕翻车，二维路线的稳定性明显优于视频驱动。",
         lead="选型建议。")
    para(doc, "三维路线保留为技术备选。纯浏览器渲染的 TalkingHead 不需要服务端显卡，"
              "如果现场只能用一台笔记本演示，这是唯一稳妥的方案。")

    # 4
    heading(doc, "可直接复用的开源底座")
    para(doc, "调研的基本判断是，这个领域已经有成熟到开箱即用的整套框架，"
              "我们没有理由自己拼装语音识别、对话模型、语音合成、口型和推流链路。"
              "下表把候选底座放在同一组维度下比较，星标数反映社区规模，最近更新时间反映维护活跃度。")
    para(doc, "表中列出的是项目简称，完整仓库地址见文末参考资料。")
    table(doc,
          ["项目", "星标", "许可证", "最近更新", "定位与要点"],
          [
              ["LiveTalking", "9.6k", "Apache-2.0", "2026-09-13",
               "实时交互流式数字人引擎，支持 ernerf、musetalk、wav2lip、Ultralight 四种模型，"
               "原生支持说话被打断，可输出 WebRTC、RTMP 与虚拟摄像头，自带形象生成页与管理页。"
               "实测环境为 Ubuntu 22.04、Python 3.12、PyTorch 2.9.1、CUDA 12.8"],
              ["Open-LLM-VTuber", "13.9k", "以仓库为准", "2026-05-15",
               "Live2D 形象加实时语音对话，支持语音打断，可完全离线运行，内置多种对话模型、语音合成与识别后端，"
               "提供网页版与桌面客户端。远程访问必须配置 https，否则浏览器不给麦克风权限"],
              ["OpenTalking", "3.1k", "Apache-2.0", "2026-09-04",
               "全链路编排框架，覆盖会话状态、对话模型、语音识别与合成、打断控制、字幕事件和 WebRTC 播放。"
               "提供无需模型权重的 mock 模式，可以在没有显卡时先打通链路"],
              ["Linly-Talker", "3.5k", "MIT", "2026-02-10",
               "集成语音识别、对话模型、语音合成与口型生成的完整系统，文档与教程最完整，"
               "适合团队用来理解整条链路"],
              ["Fay", "13.5k", "GPL-3.0", "2026-09-21",
               "数字人框架，向上适配多种数字人模型，向下接各类大语言模型，支持打断、知识库、"
               "多终端接入与表情输出。许可证是 GPL-3.0，后续涉及商用需谨慎"],
              ["Duix.Avatar", "15.6k", "以仓库为准", "2026-04-21",
               "离线数字人克隆与视频合成，支持文本与语音驱动、口型同步、多语种，提供 Docker 部署方式"],
              ["VideoChat", "1.3k", "以仓库为准", "2025-12-18",
               "实时语音交互数字人，公开给出延迟与显存实测数据，可以作为我们做延迟预算的基线"],
              ["TalkingHead", "1.6k", "以仓库为准", "2026-09-25",
               "纯浏览器端实时三维口型，可配合浏览器内的语音合成与语言模型使用，不需要服务端显卡"],
          ],
          [2.6, 1.2, 2.1, 2.0, 8.3],
          aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
                  WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER,
                  WD_ALIGN_PARAGRAPH.LEFT],
          size=8.5, keep_together=False)
    para(doc, "LiveTalking 最值得优先验证。它把实时音视频推流、口型和打断全部打包好了，"
              "自带前端页面与形象生成页，我们的工作会收缩成接对话模型、换形象、压延迟三件事。"
              "opentalking 的价值在 mock 模式，可以在还没拿到显卡、模型权重尚未下载的时候，"
              "先把对话模型到前端播放的链路打通，不阻塞进度。"
              "Open-LLM-VTuber 的产品形态和情感陪伴赛题几乎重合，适合直接作为陪伴型演示的底座，"
              "代价是形象偏桌宠风格、长期记忆功能目前暂时下线。",
         lead="推荐顺序。")

    # 5
    heading(doc, "语音链路选型")
    heading(doc, "语音识别", 2)
    para(doc, "语音识别首选 FunASR，星标 20.5k，MIT 许可，中文识别效果好，支持流式识别，"
              "2026 年 9 月 25 日仍有代码提交，国内下载顺畅。WhisperX 在词级时间戳与说话人分离上更强，"
              "但中文流式场景不如 FunASR 顺手。Qwen3-ASR 是较新的方案，可以作为对比实验，不必作为主选。")

    heading(doc, "语音合成", 2)
    para(doc, "赛题要求共情表达，所以语音合成的核心指标不是像不像真人，而是能不能带着情绪说话。"
              "下表按情感可控程度排列了几个主流开源方案。")
    table(doc,
          ["项目", "星标", "许可证", "情感控制", "关键能力"],
          [
              ["IndexTTS", "24.2k", "以仓库为准", "强",
               "IndexTTS-2.5 支持细粒度情感控制与音色情感解耦，零样本音色克隆，时长可调零点五到两倍，"
               "支持拼音与音素纠音，可用 vLLM 加速部署"],
              ["CosyVoice", "23.8k", "Apache-2.0", "强",
               "Fun-CosyVoice 3.0 支持指令式情感、语速、音量控制，双向流式合成延迟低至 150 毫秒，"
               "覆盖九种语言与十八种中文方言"],
              ["VoxCPM", "38.0k", "以仓库为准", "中高",
               "可用自然语言描述设计音色与情感，二吉字节参数，支持三十种语言与 48 千赫兹输出，"
               "音质最好但推理开销最大"],
              ["GPT-SoVITS", "62.1k", "MIT", "中",
               "一分钟语音数据即可克隆音色，社区最成熟，情感主要依靠参考音频带出"],
              ["ChatTTS", "39.9k", "AGPL-3.0", "中",
               "面向日常对话的口语化合成，语气自然，但许可证为 AGPL-3.0，商用需注意"],
              ["Fish-Speech", "32.8k", "以仓库为准", "中",
               "开源语音合成方案的先进水平，被 Duix.Avatar 采用"],
              ["edge-tts", "12.0k", "以仓库为准", "弱",
               "零成本、零显存、秒级可用，作为演示兜底"],
          ],
          [2.6, 1.3, 2.0, 2.2, 8.1],
          aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
                  WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER,
                  WD_ALIGN_PARAGRAPH.LEFT],
          size=8.5)
    para(doc, "主力方案建议 IndexTTS-2.5 或 Fun-CosyVoice 3.0，兜底方案用 edge-tts。"
              "IndexTTS 的情感控制最直接对应赛题要求，CosyVoice 的 150 毫秒流式延迟对体感帮助最大。"
              "准备两路语音合成是有必要的，演示当天任何一路出问题，另一路都能顶上。",
         lead="推荐方案。")

    heading(doc, "断句与打断", 2)
    para(doc, "语音活动检测用于判断用户是否在说话、何时说完，推荐 Silero VAD 或语音识别自带的能力。"
              "打断是指用户开口时立即停止播放、清空待播队列、丢弃当前未完成的回复，"
              "这是陪伴感和客服感的分水岭，评委也很容易感知到。"
              "LiveTalking、Open-LLM-VTuber、opentalking 都原生支持打断，这也是优先选它们的重要原因。")

    # 6
    heading(doc, "口型同步方案")
    para(doc, "口型同步有参数驱动、视频重绘和三维表情三条路线，差别集中在显存、延迟和真实感的取舍上。")
    table(doc,
          ["技术路线", "代表项目", "显存", "延迟", "真实感", "稳定性"],
          [
              ["参数与音素驱动", "TalkingHead、live2d-py、Rhubarb Lip Sync、uLipSync",
               "几乎为零", "毫秒级", "低", "高"],
              ["视频重绘驱动", "Wav2Lip、MuseTalk、LatentSync、EchoMimicV2",
               "数吉字节", "中高", "高", "中"],
              ["三维表情驱动", "Audio2Face、ARKit 表情系数", "服务端零", "低", "中", "高"],
          ],
          [2.6, 5.8, 2.0, 1.9, 1.9, 2.0],
          aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT,
                  WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER,
                  WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER])
    para(doc, "既然形象走二维路线，口型就用参数驱动，把显卡省给对话模型和语音合成。"
              "如果决赛需要展示真人形象，MuseTalk 是当前实时唇形同步的成熟选择，"
              "VideoChat 和 Linly-Talker 都采用了它。")

    # 7
    heading(doc, "与对话模型的联动架构")
    heading(doc, "数据流", 2)
    para(doc, "整条链路从用户开口到数字人回应，依次经过下面的环节。"
              "每一环都存在延迟，累加起来就是用户感受到的等待时间。")
    flow = doc.add_paragraph()
    flow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    flow.paragraph_format.space_before = Pt(4)
    flow.paragraph_format.space_after = Pt(10)
    flow.paragraph_format.line_spacing = 1.5
    set_run(flow.add_run("麦克风　→　静音检测　→　语音识别　→　对话模型　→　分句　→　语音合成"
                         "　→　口型驱动　→　WebRTC 推流　→　浏览器播放与字幕同步"),
            size=10, bold=True)

    heading(doc, "协议选型", 2)
    para(doc, "音视频推流用 WebRTC，它延迟最低，LiveTalking、opentalking 和 Linly-Talker 的流式版本都采用它，"
              "代价是需要放通 UDP 的大范围端口。对话模型调用走 HTTP 加流式返回，"
              "与张力文约定统一的接口形式，他们后续换模型时我这里不需要改代码。"
              "控制信令与字幕事件走 WebSocket，用于发送打断指令、会话状态和逐字字幕。")

    heading(doc, "三个必须做对的设计", 2)
    para(doc, "不要等大模型把整段话说完才送去合成。一旦遇到句号、问号或叹号就切句送合成，"
              "首包延迟可以从三秒压到一点五秒以内。",
         lead="流式分句。")
    para(doc, "用户一开口就立刻停播、清空队列。实现上要让语音活动检测的事件优先于待播队列。",
         lead="打断优先。")
    para(doc, "对话模型除了返回文本，还要返回情绪标签，这个标签同时驱动语音合成的情感风格与语速音量，"
              "以及数字人的表情和动作。效果是用户倾诉低落时数字人语气变温柔、表情收敛，"
              "说到开心的事时语气轻快、表情明亮。这一步是我们和套壳数字人方案的分界线。",
         lead="情感联动。")

    heading(doc, "延迟预算", 2)
    para(doc, "目标是把首包控制在一点五秒以内，最差不超过三秒。各环节的预算分配如下。")
    table(doc,
          ["环节", "目标耗时", "说明"],
          [
              ["静音检测判定用户说完", "0.1 秒", "端点检测阈值可调，设得太紧会抢话"],
              ["语音识别", "0.3 秒", "流式识别，边说边出字"],
              ["对话模型首 token", "0.3 至 1.0 秒", "取决于模型规模，小模型优先"],
              ["首句分句与语音合成首包", "0.15 至 0.5 秒", "CosyVoice 双向流式可低至 150 毫秒"],
              ["口型生成与推流", "0.2 至 0.5 秒", "参数驱动远快于视频重绘"],
              ["网络往返", "0.1 秒", "本地演示可忽略"],
              ["合计", "1.15 至 2.5 秒", "控制在三秒内即达到同行水平"],
          ],
          [5.4, 3.2, 7.6],
          aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
          bold_rows=(6,))

    heading(doc, "部署注意事项", 2)
    bullet(doc, "浏览器只在 https 或 localhost 下允许访问麦克风，远程访问必须配置 https 反向代理。")
    bullet(doc, "WebRTC 需要放通 UDP 大范围端口，校园网和云服务器的安全组要提前配置。")
    bullet(doc, "模型权重优先从 ModelScope 获取，HuggingFace 直连大概率失败。")
    bullet(doc, "级联方案约需八吉字节显存，单张 4090 余量充足，但不要在同一张卡上开多路并发。")
    bullet(doc, "Fay 是 GPL-3.0、ChatTTS 是 AGPL-3.0，若项目后续有商用规划，选型时要避开或单独评估。")

    # 8
    heading(doc, "情绪标签体系")
    para(doc, "这套标签是对话模型、语音合成和数字人之间的公共接口，需要张力文和唐应杰一起确认，"
              "否则情感联动无法实现，评估口径也没法和交互表现对齐。下面是六类标签的草案。")
    table(doc,
          ["标签", "典型场景", "语音合成的情感映射", "数字人表现"],
          [
              ["开心", "分享好消息、被认可", "轻快、语速略快、音调上扬", "微笑、点头"],
              ["平静", "日常闲聊", "自然、中速", "平稳、偶尔眨眼"],
              ["低落", "情绪倾诉、受挫", "温柔、语速放慢、音量略降", "表情收敛、倾听姿态"],
              ["焦虑", "考试、就业压力", "沉稳、语速放缓、语气笃定", "身体前倾、安抚性点头"],
              ["愤怒", "抱怨、人际冲突", "平静克制、语速平缓", "中性表情、避免过度反应"],
              ["惊喜", "意外之喜", "明亮、适度上扬", "睁大眼睛、轻微后仰"],
          ],
          [1.8, 4.2, 5.0, 5.2],
          aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT,
                  WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])
    para(doc, "有两处需要注意。标签数量不宜过多，否则语音合成的情感参数难以区分，数据标注成本也会失控。"
              "如果唐应杰的评估口径里包含回复情感恰当性这一项，这套标签可以直接作为对齐依据。")

    # 9
    heading(doc, "演示方案构想")
    heading(doc, "场景设定", 2)
    para(doc, "演示的场景定位是面向高校学生的情绪陪伴数字人。用户是压力较大的大学生，"
              "数字人是可信赖的同龄伙伴。可演示的痛点有三个，分别是学业受挫、就业焦虑和宿舍人际矛盾，"
              "其中学业受挫最容易在短时间里让观众产生代入感。")

    heading(doc, "方案分档", 2)
    para(doc, "用 Open-LLM-VTuber 或 LiveTalking 加上 Live2D 形象、"
              "CosyVoice 或 IndexTTS，再接上对话模型服务，一到两周可以完成。这一档保证团队有可演示的成品，建议先做。",
         lead="稳妥版。")
    para(doc, "在稳妥版基础上换用 opentalking 或 LiveTalking 的编排能力，"
              "叠加真人二维形象、情感联动和打断，需要三到四周。这一档作为决赛主推方案。",
         lead="进阶版。")
    para(doc, "端到端语音对话模型配合数字人，显存需求约二十吉字节、首包约七秒。"
              "体验明显更差、资源明显更贵，只作为技术展望，不作为主方案。",
         lead="理想版。")

    heading(doc, "三分钟演示脚本", 2)
    table(doc,
          ["时间", "演示内容", "展示的技术点"],
          [
              ["0:00 至 0:40", "用户语音提问，说这次考试又没过，感觉特别糟糕",
               "完整语音链路，静音检测、语音识别与对话模型的共情回应"],
              ["0:40 至 1:20", "数字人温柔回应，语速放慢、表情收敛", "情感联动，情绪标签驱动语音与表情"],
              ["1:20 至 1:50", "用户中途打断，说等一下我先说个别的事", "打断能力，数字人立即停口倾听"],
              ["1:50 至 2:30", "用户提到上次聊过的话题，数字人主动承接", "记忆融合，与对话模型模块的联合成果"],
              ["2:30 至 3:00", "展示界面上的情绪曲线与会话记录", "用可视化证明共情效果，而非主观感受"],
          ],
          [2.6, 6.6, 7.0],
          aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])
    para(doc, "脚本需要提前和张力文、唐应杰对齐。第三分钟用到的记忆与情绪能力是他们模块的产出，"
              "接口和演示用例必须提前约定，不能到现场才拼接。")

    heading(doc, "降级预案", 2)
    bullet(doc, "现场网络差或云端服务不可用，改为全部本地部署，并准备提前录制的视频兜底。")
    bullet(doc, "显卡不可用，切换到纯浏览器渲染方案，或直接使用预渲染视频演示。")
    bullet(doc, "语音合成太慢或卡顿，切换到 edge-tts，零显存且秒级可用。")
    bullet(doc, "麦克风权限被浏览器拦截，改为手动输入文字驱动，功能展示不受影响。")
    bullet(doc, "推流掉帧，降低输出分辨率与帧率，优先保证口型与声音同步。")

    # 10
    heading(doc, "风险与对策")
    table(doc,
          ["风险", "影响", "对策"],
          [
              ["本模块对初赛排名没有贡献，却占用时间", "高",
               "国庆前只做调研与选型，动手实现放在初赛提交之后"],
              ["数字人服务影响初赛镜像稳定性", "高", "数字人独立成服务，初赛镜像不集成"],
              ["开发算力被演示环境占用", "中", "训练与演示分时使用，演示优先选参数驱动和轻量语音合成"],
              ["情绪标签在三方之间不统一", "中", "把情绪标签草案尽早提交三方确认"],
              ["开源项目停止维护或许可证冲突", "中", "主方案避开 AGPL 与 GPL 项目，锁定版本号，不追主干最新提交"],
              ["决赛现场演示翻车", "高", "准备五条降级预案，并提前录制演示视频"],
          ],
          [5.6, 1.6, 9.0],
          aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])

    # 11
    heading(doc, "国庆前行动清单")
    table(doc,
          ["时间", "动作", "产出", "配合人"],
          [
              ["9 月 25 日至 26 日",
               "用 opentalking 的 mock 模式或 edge-tts 跑通文本到语音到网页播放的最小链路，不依赖显卡与模型权重",
               "一段可播放的演示音频与链路截图", "无"],
              ["9 月 27 日至 28 日",
               "在咪咕仝学实例上验证 LiveTalking 或 Open-LLM-VTuber 能否跑起来",
               "跑通与否的结论，以及显存和延迟的实测数字", "邹艾融　确认算力领取"],
              ["9 月 29 日",
               "定稿本笔记，补上实测数据，并提出情绪标签体系",
               "调研笔记定稿", "无"],
              ["9 月 30 日",
               "与钟老师讨论数字人与初赛镜像是否分离、情感联动是否作为项目创新点、决赛演示形式",
               "讨论纪要", "全队"],
          ],
          [2.8, 6.4, 4.2, 2.8],
          aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT,
                  WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

    # 12
    heading(doc, "参考资料")
    para(doc, "实时对话底座：LiveTalking https://github.com/lipku/LiveTalking；"
              "Open-LLM-VTuber https://github.com/Open-LLM-VTuber/Open-LLM-VTuber；"
              "OpenTalking https://github.com/datascale-ai/opentalking；"
              "Linly-Talker https://github.com/Kedreamix/Linly-Talker；"
              "Fay https://github.com/xszyou/Fay；"
              "Duix.Avatar https://github.com/duixcom/Duix-Avatar；"
              "VideoChat https://github.com/Henry-23/VideoChat。",
         size=9, after=6)
    para(doc, "口型同步：MuseTalk https://github.com/TMElyralab/MuseTalk；"
              "Wav2Lip https://github.com/Rudrabha/Wav2Lip；"
              "LatentSync https://github.com/bytedance/LatentSync；"
              "EchoMimicV2 https://github.com/antgroup/echomimic_v2；"
              "TalkingHead https://github.com/met4citizen/TalkingHead。",
         size=9, after=6)
    para(doc, "语音：FunASR https://github.com/modelscope/FunASR；"
              "IndexTTS https://github.com/index-tts/index-tts；"
              "CosyVoice https://github.com/QwenAudio/CosyVoice；"
              "VoxCPM https://github.com/OpenBMB/VoxCPM；"
              "GPT-SoVITS https://github.com/RVC-Boss/GPT-SoVITS；"
              "edge-tts https://github.com/rany2/edge-tts。",
         size=9, after=6)
    para(doc, "赛事资料：2026 动感地带 AI 加高校创智计划赛事介绍，第 11 至 15 页为技术赛道内容；"
              "技术赛道的工程包与数据集下载地址见赛事介绍第 12 页。",
         size=9, after=6)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    props = doc.core_properties
    props.author = "郭柱"
    props.last_modified_by = "郭柱"
    props.title = "数字人交互模块技术调研与演示方案"
    props.subject = "2026 动感地带 AI 高校创智计划 技术赛道 数字人交互模块"
    doc.save(OUT)
    print("saved:", OUT)


if __name__ == "__main__":
    build()
