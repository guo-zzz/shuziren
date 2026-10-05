# -*- coding: utf-8 -*-
"""Build the dual-persona + avatar asset DOCX (deliverable for the AI track)."""

import os
import sys

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_docx as B  # 复用上一份文档的排版原语

OUT = r"D:\iwen-codex\codex-2\deliverables\郭柱-双分人人设与形象素材说明.docx"
ASSETS = r"D:\iwen-codex\codex-2\demo\assets"
OVERVIEW = os.path.join(ASSETS, "形象素材总览.png")


def quote_block(doc, lines, fill="F7F8FA", size=9):
    """单格表格当作引文/代码块，避免长 prompt 排版散掉。"""
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    B._borders(t)
    cell = t.rows[0].cells[0]
    cell.width = Cm(16.2)
    B._shade(cell, fill)
    cell.text = ""
    first = True
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        pf = p.paragraph_format
        pf.space_after = Pt(2)
        pf.space_before = Pt(2)
        pf.line_spacing = 1.3
        B.set_run(p.add_run(line), size=size)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(doc, path, width_cm, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(path, width=Cm(width_cm))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(10)
    B.set_run(c.add_run(caption), size=9, color="5A6472")


def build():
    doc = Document()
    B.setup_styles(doc)

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    B.footer_page_number(section)

    # ---------------------------------------------------------------- 标题
    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(2)
    B.set_run(title.add_run("双分人人设与形象素材说明"), size=21, bold=True, font=B.HEAD_FONT)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(14)
    B.set_run(
        sub.add_run("2026 动感地带 AI+ 高校创智计划 · 技术赛道 · 数字人综合情感陪伴对话模型"
                    "｜编制：郭柱｜2026 年 10 月 2 日"),
        size=9.5, color="5A6472")

    B.para(doc,
           "这份文档交付两件事：一是「当下的我」和「理想的我」两套人设的完整定义，包括它们各自的"
           "定位、说话方式、systemPrompt 原文和禁止表达；二是与两套人设配套的 14 个数字人形象素材，"
           "以及素材在 PPT、Demo 页面里的使用方式。人和形象都放在同一个工程里，"
           "改一处文案或配色就能同步生效，不需要重新画图。",
           after=10)
    B.para(doc,
           "文档面向三类读者：后端同学照人设写 prompt 与情绪标签；测试同学照人设判断回复是否符合预期；"
           "答辩演示时照第五节和第七节取用对照表与形象素材。"
           "配套的前端工程、启动方式和接口细节写在工程目录的 README.md 里，本文不重复。",
           after=12)

    # ---------------------------------------------------------- 一、交付内容
    B.heading(doc, "一、交付内容与文件位置")
    B.para(doc,
           "本阶段的交付物是「双分人人设与形象素材」加「前端简易 Demo 工程」。"
           "其中人设不是一份静态文案，而是直接驱动程序运行的配置文件；"
           "形象也不是图片素材，而是用矢量图形代码画出来的，因此改配色即换形象，"
           "并且可以随时导出成 PPT 能用的素材。",
           after=8)
    B.table(
        doc,
        ["交付内容", "文件位置", "说明与用途", "主要使用者"],
        [
            ["两套人设定义", "demo/web/personas.js",
             "人设文案、说话风格、systemPrompt、安全边界、情绪词表；程序直接读取", "张力文、王如怡"],
            ["形象绘制与动画", "demo/web/avatar.js",
             "数字人的矢量绘图与表情、口型动画；表情强度参数在文件开头", "郭柱"],
            ["形象素材（矢量）", "demo/assets/svg/",
             "14 个 SVG，2 套人设 × 6 种情绪加 1 个说话状态；可无损放大", "全队、答辩 PPT"],
            ["形象素材（位图）", "demo/assets/png/",
             "同款 14 个 PNG，480×620，兼容不支持 SVG 的旧版 Office", "全队、答辩 PPT"],
            ["素材总览图", "demo/assets/形象素材总览.png",
             "一页看完全部形象与情绪差异，可直接贴进汇报材料", "邹艾融、王如怡"],
            ["前端工程与说明", "demo/ 与 demo/README.md",
             "启动方式、后端接入方法、情绪标签协议、演示流程", "全队"],
        ],
        [3.0, 4.3, 6.5, 2.4],
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT,
                WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])
    B.para(doc,
           "素材的生成脚本是 demo/tools/export-avatars.mjs，改完人设或配色后重新执行一次，"
           "assets 目录下的 SVG、清单和总览页会整体刷新，不会出现图文不一致。",
           after=12, size=9.5)

    # -------------------------------------------------------- 二、设计依据
    B.heading(doc, "二、双分身的设计依据")
    B.para(doc,
           "两个分身不是把同一个回答换个语气说两遍。它们的依据是心理学中的自我差异理论"
           "（Higgins, 1987）：人的自我分为「现实自我」和「理想自我」，两者之间的落差会带来情绪压力，"
           "而把落差看清楚、并知道下一步能做什么，落差就不再只是压力。",
           after=8)
    B.para(doc,
           "落到产品上，就是让两个分身同时在场、一起回答同一个问题：左边先承认用户此刻的处境，"
           "让情绪落地；右边再给出一个今天就能做的小动作，让落差变成可以着手的方向。"
           "两个分身共用同一段用户输入，因此观点必须是同一件事的两个视角，而不是两种人格互相拆台。"
           "这也是本作品回应命题的方式：情感陪伴的重点不在陪聊，而在于情绪被接住之后还能往前走一步。",
           after=12)

    # ------------------------------------------------------ 三、当下的我
    B.heading(doc, "三、当下的我（左，present）")
    B.para(doc,
           "这个分身回答的是「我现在这样是不是也说得过去」。它的任务是接住此刻真实的自己，"
           "先承认困难的存在，再陪用户把话说清楚。它不催促、不评判、不把问题转回用户身上，"
           "也不用励志话术把情绪盖过去。",
           after=8)
    B.table(
        doc,
        ["项目", "设定"],
        [
            ["角色定位", "接住此刻真实的自己"],
            ["一句话气质", "不评判、不催促，先承认这件事就是很难"],
            ["语气特征", "温和、具体、用短句；多用「我能理解」「这确实很难」这类承接句"],
            ["回应长度", "60 到 120 字，一问一答节奏略慢"],
            ["主色", "#5A6B8C（沉稳的蓝灰，视觉上更安静）"],
            ["禁止表达", "「加油就好了」「别想太多」「你应该早点……」等催促、评判式说法"],
        ],
        [3.6, 12.6],
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])
    B.para(doc, "systemPrompt 原文（后端直接使用，节选自 personas.js）：", after=4, size=9.5)
    quote_block(doc, [
        "你是「当下的我」，用户内心那个只会接住情绪、不催促的部分。",
        "说话方式：温和、具体、不评判；先承认这件事确实很难，再陪用户把事情说清楚。",
        "不要说「加油就好了」「别想太多」；不要给三步五点的大计划。",
        "回复控制在 60 到 120 字，像面对面说话那样自然。",
        "每次回复的第一行先给出情绪词（开心、平静、低落、焦虑、愤怒、惊喜 中的一个），第二行起是正文。",
    ])
    if os.path.exists(os.path.join(ASSETS, "png", "present-pingjing.png")):
        figure(doc, os.path.join(ASSETS, "png", "present-pingjing.png"), 5.6,
               "图 1　当下的我（平静表情）")

    # ------------------------------------------------------ 四、理想的我
    B.heading(doc, "四、理想的我（右，ideal）")
    B.para(doc,
           "这个分身回答的是「我可以往哪儿去」。它不重复安慰，而是给出一个足够小、今天就能开始的步骤，"
           "并用自己的经验说话，让建议听起来是走过来的人的感受，而不是要求。"
           "它必须避开「你怎么还没做到」这类隐含指责的句式，否则两个分身会一起压住用户。",
           after=8)
    B.table(
        doc,
        ["项目", "设定"],
        [
            ["角色定位", "看见可以抵达的样子"],
            ["一句话气质", "笃定、明亮，把落差拆成今天能做的一小步"],
            ["语气特征", "用第一人称经验说话，句子干净，结尾落到具体动作或具体感受"],
            ["回应长度", "60 到 120 字，与左边的长度相当，便于双栏对齐"],
            ["主色", "#C08428（赭金色，比左边更明亮、更有行动感）"],
            ["禁止表达", "「你怎么还没做到」「别人都行」等比较、指责式说法"],
        ],
        [3.6, 12.6],
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])
    B.para(doc, "systemPrompt 原文（后端直接使用，节选自 personas.js）：", after=4, size=9.5)
    quote_block(doc, [
        "你是「理想的我」，用户内心那个已经走过这段路、知道方向的部分。",
        "说话方式：笃定、明亮、有行动感；用第一人称经验说话，给一个今天就能做的小步骤。",
        "不要说「你怎么还没做到」「别人都行」；不要否定用户此刻的感受。",
        "回复控制在 60 到 120 字，和「当下的我」保持相近的长度。",
        "每次回复的第一行先给出情绪词（开心、平静、低落、焦虑、愤怒、惊喜 中的一个），第二行起是正文。",
    ])
    if os.path.exists(os.path.join(ASSETS, "png", "ideal-pingjing.png")):
        figure(doc, os.path.join(ASSETS, "png", "ideal-pingjing.png"), 5.6,
               "图 2　理想的我（平静表情）")

    # ------------------------------------------------------ 五、差异对照
    B.heading(doc, "五、两个分身的差异对照")
    B.para(doc,
           "演示和测试时按这张表判断两个分身是否各司其职。最常出现的问题是两边都在安慰，"
           "或者两边都在给建议，出现任一种就说明 prompt 或知识库内容带偏了。",
           after=8)
    B.table(
        doc,
        ["对照维度", "当下的我", "理想的我"],
        [
            ["回答的问题", "我现在这样，是不是也说得过去", "我可以往哪儿去"],
            ["情绪姿态", "先承接，允许情绪停留", "先看见方向，再把落差拆小"],
            ["开场白", "「嗯，你来了。想说什么都可以，我在这儿听着。」",
             "「我在这儿。我们不着急，先聊聊今天发生了什么。」"],
            ["结尾动作", "把用户的感受复述清楚，确认被听懂", "给一个今天就能做的小步骤"],
            ["语言特征", "短句、温和、多用承接句", "第一人称经验、句子干净、落到动作"],
            ["回复长度", "60 到 120 字", "60 到 120 字"],
            ["主色", "#5A6B8C", "#C08428"],
            ["禁止表达", "催促与评判（加油就好了、别想太多）", "比较与指责（你怎么还没做到）"],
        ],
        [3.2, 6.5, 6.5],
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])
    B.para(doc,
           "两个分身默认按顺序开口，间隔 0.7 秒，左边先说、右边后说。"
           "这个顺序让「先接住、再往前」的层次在演示时听得出来；"
           "如果答辩现场希望制造同时回应的效果，可以在设置里关掉顺序发言。",
           after=12, size=9.5)

    # ------------------------------------------------------ 六、情绪与表情
    B.heading(doc, "六、情绪词与表情参数")
    B.para(doc,
           "数字人的表情由情绪词驱动，情绪词同时决定语气和表情，这是本作品「情感联动」的具体做法。"
           "可用情绪词固定为六个，后端不要自造词；参数写在 demo/web/avatar.js 开头的 EMOTION_PRESETS 里，"
           "调整这些数值就能改变表情强度，不需要改绘图代码。",
           after=8)
    B.table(
        doc,
        ["情绪词", "典型场景", "嘴角弧度", "眉毛角度", "眼睛开合", "腮红浓度"],
        [
            ["开心", "有进展、被夸奖、好消息", "+1.0（明显上翘）", "上挑 3 度", "1.00", "0.55"],
            ["平静", "日常闲聊、开场、倾听", "+0.25（轻微上扬）", "水平", "1.00", "0.22"],
            ["低落", "受挫、失望、被拒绝", "-0.80（下撇）", "下压 7 度", "0.86（略微垂眼）", "0.12"],
            ["焦虑", "担心、紧张、反复犹豫", "-0.25（略微紧绷）", "下压 4 度", "1.02", "0.20"],
            ["愤怒", "被冒犯、强烈不满", "-0.90（紧抿）", "内压 11 度", "0.92", "0.28"],
            ["惊喜", "意外的好事、被理解", "+0.50（张开笑）", "上抬 5 度", "1.16（睁大）", "0.50"],
        ],
        [1.8, 4.4, 3.4, 2.6, 2.6, 1.8],
        aligns=[WD_ALIGN_PARAGRAPH.CENTER] + [WD_ALIGN_PARAGRAPH.LEFT] * 2
               + [WD_ALIGN_PARAGRAPH.CENTER] * 3)
    B.para(doc,
           "说话时的口型不是预设动画，而是按文字出现的节奏实时开合：接上语音合成之后，"
           "改调一个接口就能换成用真实音量驱动，这一点已经预留好了。",
           after=12, size=9.5)

    # ------------------------------------------------------ 七、素材清单
    B.heading(doc, "七、形象素材清单与使用方式")
    B.para(doc,
           "素材共 14 个：两套人设各 6 种情绪，再加 1 个说话中的状态。"
           "每个素材同时提供 SVG 与 PNG 两种格式，矢量版放进 PPT 可以任意放大不糊，"
           "位图版用于不支持 SVG 的旧版本 Office。下面的总览图同时展示了全部素材，"
           "可以直接复制进汇报材料。",
           after=10)
    if os.path.exists(OVERVIEW):
        figure(doc, OVERVIEW, 16.2, "图 3　数字人形象素材总览（上排：当下的我；下排：理想的我）")

    B.para(doc, "文件命名规则：", after=4, size=9.5)
    quote_block(doc, [
        "present-kaixin.svg     当下的我 · 开心          ideal-kaixin.svg     理想的我 · 开心",
        "present-pingjing.svg   当下的我 · 平静          ideal-pingjing.svg   理想的我 · 平静",
        "present-diluo.svg      当下的我 · 低落          ideal-diluo.svg      理想的我 · 低落",
        "present-jiaolv.svg     当下的我 · 焦虑          ideal-jiaolv.svg     理想的我 · 焦虑",
        "present-fennu.svg      当下的我 · 愤怒          ideal-fennu.svg      理想的我 · 愤怒",
        "present-jingxi.svg     当下的我 · 惊喜          ideal-jingxi.svg     理想的我 · 惊喜",
        "present-speaking.svg   当下的我 · 说话中        ideal-speaking.svg   理想的我 · 说话中",
    ], fill="F7F8FA", size=8.5)

    B.para(doc, "使用方式：", after=4)
    B.bullet(doc, "插入 PPT：直接拖入 assets/svg 里的文件即可，缩放不失真；需要重新着色时可在 PowerPoint 里取消组合。")
    B.bullet(doc, "预览某个表情：页面地址后加 ?pose=低落，两个分身会定格在该表情；加 ?speak=1 保持说话口型。")
    B.bullet(doc, "单独导出：Demo 页面每个分身标题右侧有「导出形象」按钮，会按当前表情存成一个 SVG 文件。")
    B.bullet(doc, "总览页：启动 Demo 后访问 http://127.0.0.1:8090/assets/contact-sheet.html，"
                  "可在浏览器里放大查看每一个形象。")

    # ------------------------------------------------------ 八、接口约定
    B.heading(doc, "八、与其他成员的接口约定")
    B.para(doc,
           "前端按 persona_id 分别为 present 和 ideal 发起请求，各自带上自己的 systemPrompt。"
           "后端不需要额外分支逻辑，按传进来的 systemPrompt 回答即可；"
           "如果接入了心理学 skill 知识库，把知识库内容拼在 systemPrompt 之后，"
           "不要删掉原有的角色定位和安全边界。",
           after=8)
    B.para(doc,
           "情绪标签是前后端唯一需要额外约定的字段，两种给法任选：在流式返回的 delta 里带 emotion 字段，"
           "或者让回复的第一行是方括号包起来的情绪词（例如「【低落】」），第二行起为正文。"
           "两种都不给时按「平静」处理，不会报错。",
           after=8)
    B.para(doc,
           "安全边界写在 personas.js 的 SAFETY_RULES 里，两个分身共用：不做诊断、不下医学结论、"
           "遇到自伤或自杀等危机信号立刻停止角色化回应，建议联系学校心理中心、辅导员或拨打全国心理援助热线 12356。"
           "这段内容已经拼进两个 systemPrompt，后端在拼接知识库时应保留。",
           after=12)

    # ------------------------------------------------------ 九、边界与下一步
    B.heading(doc, "九、尚未完成的部分与下一轮建议")
    B.para(doc,
           "本阶段的目标是先把核心效果跑通，因此语音、形象精度和适配都刻意做了简化，"
           "这些不是遗漏，而是排期结果。",
           after=8)
    B.table(
        doc,
        ["项目", "当前状态", "下一轮建议"],
        [
            ["语音", "未接入语音合成，口型由文字节奏驱动",
             "接入 IndexTTS 或 CosyVoice，再用真实音量驱动口型（接口已预留）"],
            ["形象精度", "简易卡通矢量形象",
             "如需照片级，准备正面人像素材，在嘴部锚点做局部变形"],
            ["历史与会话", "刷新页面即清空",
             "接入本地存储或后端会话，支持多轮长对话回看"],
            ["适配", "按桌面浏览器设计",
             "演示若用平板或投影，补一版窄屏布局"],
        ],
        [2.6, 5.6, 8.0],
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])
    B.para(doc,
           "优先级建议：语音与真实口型联动最能提升答辩观感，建议排在下一轮第一位；"
           "形象精度提升的收益取决于答辩现场是否需要特写，排在语音之后。",
           after=10)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    props = doc.core_properties
    props.author = "郭柱"
    props.last_modified_by = "郭柱"
    props.title = "双分人人设与形象素材说明"
    props.subject = "2026 动感地带 AI+ 高校创智计划 技术赛道 数字人综合情感陪伴对话模型"
    doc.save(OUT)
    print("saved:", OUT)


if __name__ == "__main__":
    build()